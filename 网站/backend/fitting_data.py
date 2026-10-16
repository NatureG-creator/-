"""有界表格读取和显式行审计；不修改原上传数据。"""
from __future__ import annotations

import csv
import io
import math
from pathlib import Path

MAX_ROWS, MAX_COLUMNS, MAX_BYTES = 5000, 64, 4_000_000


def _cell(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    try:
        number = float(value)
        return number if math.isfinite(number) else str(value)
    except (ValueError, TypeError):
        return str(value).strip()[:1024] or None


def _table(raw, **metadata):
    rows, category_rows = [], []
    for row in raw:
        row = list(row)
        if len(row) > MAX_COLUMNS:
            raise ValueError('表格最多64列')
        if not row or all(v is None or str(v).strip()=='' for v in row):
            if not rows:continue
            row=[None]*len(rows[0])
        rows.append([_cell(v) for v in row])
        category_rows.append([None if v is None else str(v).strip() for v in row])
        if len(rows) > MAX_ROWS+1:
            raise ValueError('表格最多5000行数据')
    if not rows:
        raise ValueError('表格没有数据')
    while rows and all(value is None for value in rows[-1]):
        rows.pop();category_rows.pop()
    if not rows:raise ValueError('表格没有数据')
    if len({len(row) for row in rows}) != 1:
        raise ValueError('表格各行列数不一致，请检查分隔符和空列')
    header = all(isinstance(v,str) and v for v in rows[0])
    if header:
        names = rows.pop(0)
        category_rows.pop(0)
    else:
        names = [f'列{i+1}' for i in range(len(rows[0]))]
    if len(rows)>MAX_ROWS:
        raise ValueError('表格最多5000行数据')
    if not rows or len(names)<2:
        raise ValueError('至少需要两列和一行数据')
    # 列用索引选择，即使重名也不会错选；界面同时显示编号。
    return dict(columns=names,rows=rows,category_rows=category_rows,row_count=len(rows),header=header,**metadata)


def load_table(path=None, text=None, sheet=None, variable=None):
    if (path is None)==(text is None):
        raise ValueError('请提供一个上传文件或粘贴文本')
    metadata=dict(sheets=[],variables=[],sheet=None,variable=None)
    if path is not None:
        path=Path(path)
        if path.stat().st_size>MAX_BYTES:
            raise ValueError('拟合数据文件最多4MB；大矩阵请先导出所需列')
        suffix=path.suffix.lower()
        if suffix=='.xlsx':
            from openpyxl import load_workbook
            import zipfile
            from xml.etree.ElementTree import ParseError
            try:
                with zipfile.ZipFile(path) as archive:
                    if sum(item.file_size for item in archive.infolist())>64_000_000:
                        raise ValueError('Excel展开内容超过64MB限制')
                book=load_workbook(path,read_only=True,data_only=True)
            except (zipfile.BadZipFile,KeyError,ParseError,OSError) as exc:
                raise ValueError('XLSX格式损坏或缺少工作簿内容，请重新导出') from exc
            try:
                metadata['sheets']=book.sheetnames
                chosen=sheet or book.sheetnames[0]
                if chosen not in book.sheetnames:raise ValueError('所选工作表不存在')
                metadata['sheet']=chosen
                return _table(book[chosen].iter_rows(values_only=True),**metadata)
            except (zipfile.BadZipFile,KeyError,ParseError,OSError) as exc:
                raise ValueError('XLSX工作表内容损坏，请重新导出') from exc
            finally:book.close()
        if suffix=='.mat':
            from scipy.io import whosmat,loadmat
            from scipy.io.matlab import MatReadError
            try:
                entries=whosmat(path)
            except (NotImplementedError,ValueError,OSError,IndexError,TypeError,MatReadError) as exc:
                raise ValueError('请使用MATLAB v7 MAT（当前不读取v7.3），或导出CSV') from exc
            entries=[(name,shape) for name,shape,kind in entries
                     if len(shape)==2 and kind in {'double','single','int8','uint8','int16','uint16','int32','uint32','int64','uint64'}]
            metadata['variables']=[name for name,_ in entries]
            if not variable and len(entries)!=1:
                return dict(columns=[],rows=[],row_count=0,header=False,**metadata)
            chosen=variable or entries[0][0]
            shape=next((shape for name,shape in entries if name==chosen),None)
            if shape is None:raise ValueError('所选MAT变量不是可读取的二维实数矩阵')
            if shape[0]>MAX_ROWS or shape[1]>MAX_COLUMNS:raise ValueError('MAT矩阵超过5000行或64列限制')
            metadata['variable']=chosen
            try:array=loadmat(path,variable_names=[chosen])[chosen]
            except (ValueError,OSError,IndexError,TypeError,MatReadError) as exc:raise ValueError('MAT变量内容损坏，请重新导出') from exc
            import numpy as np
            if np.iscomplexobj(array):raise ValueError('拟合不接受复数矩阵')
            return _table(array.tolist(),**metadata)
        if suffix not in {'.csv','.txt','.tsv'}:raise ValueError('拟合支持CSV/TXT/TSV/XLSX/MAT；旧版XLS请另存XLSX')
        try:text=path.read_text(encoding='utf-8-sig')
        except UnicodeDecodeError:text=path.read_text(encoding='gb18030')
    if not isinstance(text,str) or len(text.encode('utf-8'))>MAX_BYTES:
        raise ValueError('粘贴文本最多4MB')
    sample='\n'.join(text.splitlines()[:10])
    delimiter='\t' if '\t' in sample else ',' if ',' in sample else ';' if ';' in sample else None
    raw=csv.reader(io.StringIO(text),delimiter=delimiter) if delimiter else (line.split() for line in text.splitlines())
    return _table(raw,**metadata)


def _index(selection,key,count,required=False):
    value=selection.get(key)
    if value is None and not required:return None
    if type(value) is not int or not 0<=value<count:raise ValueError(f'{key}列索引无效，请重新选择列')
    return value


def _finite(value):
    return type(value) in {float,int} and math.isfinite(value)


def prepare_dataset(table,selection):
    if not isinstance(selection,dict):raise ValueError('selection必须是对象')
    count=len(table['columns'])
    xi=_index(selection,'x_column',count,True);yi=_index(selection,'y_column',count,True)
    if xi==yi:raise ValueError('X与Y必须选择不同列')
    gi=_index(selection,'group_column',count);si=_index(selection,'sigma_column',count)
    vi=_index(selection,'validation_group_column',count)
    policy=selection.get('missing_policy','reject')
    if policy not in {'reject','drop'}:raise ValueError('缺失值策略无效')
    excluded=selection.get('exclude_rows',[])
    if not isinstance(excluded,list) or any(type(v)is not int or not 1<=v<=len(table['rows']) for v in excluded):
        raise ValueError('排除行必须是有效的1起始数据行号（不含表头）')
    low,high=selection.get('x_min'),selection.get('x_max')
    if any(v is not None and not _finite(v) for v in (low,high)) or low is not None and high is not None and low>=high:
        raise ValueError('拟合X范围必须是有限数且下界小于上界')
    data=dict(x=[],y=[],groups=[],sigma=[] if si is not None else None,
              validation_groups=[] if vi is not None else None,row_ids=[])
    audit=dict(original_rows=len(table['rows']),excluded_rows=sorted(set(excluded)),invalid_rows=[],outside_range_rows=[])
    for i,row in enumerate(table['rows'],1):
        categories=table.get('category_rows',table['rows'])[i-1]
        if i in excluded:continue
        invalid=not _finite(row[xi]) or not _finite(row[yi]) or si is not None and (not _finite(row[si]) or row[si]<=0)
        invalid=invalid or any(j is not None and row[j] is None for j in (gi,vi))
        if invalid:
            audit['invalid_rows'].append(i)
            if policy=='reject':raise ValueError(f'第{i}行X/Y不是有限数、分类缺失或测量误差非正；可选择显式剔除并记录')
            continue
        if low is not None and row[xi]<low or high is not None and row[xi]>high:
            audit['outside_range_rows'].append(i);continue
        data['x'].append(row[xi]);data['y'].append(row[yi]);data['row_ids'].append(i)
        data['groups'].append(str(categories[gi]) if gi is not None else '全部数据')
        if si is not None:data['sigma'].append(row[si])
        if vi is not None:data['validation_groups'].append(str(categories[vi]))
    if len(data['x'])<3:raise ValueError('清理和范围筛选后至少需要3个有效数据点')
    if len(set(data['groups']))>10:raise ValueError('每次最多10个分类组')
    if any(len(set(x for x,g in zip(data['x'],data['groups']) if g==label))<2 for label in set(data['groups'])):
        raise ValueError('每个分类组至少需要两个不同X值')
    audit['retained_rows']=len(data['x']);audit['duplicate_x_count']=len(data['x'])-len(set(zip(data['groups'],data['x'])))
    data['audit']=audit
    for key,fallback in (('x_label',table['columns'][xi]),('y_label',table['columns'][yi]),('x_unit',''),('y_unit','')):
        value=selection.get(key) or fallback
        if not isinstance(value,str) or len(value)>120:raise ValueError('坐标名称和单位必须是120字以内文本')
        data[key]=value
    return data
